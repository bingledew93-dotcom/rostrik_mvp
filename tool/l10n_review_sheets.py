"""Builds one translation-review spreadsheet per language, for native speakers.

Each workbook lists every user-facing string — the Flutter ARB catalogues, the
Android native strings (alarm screen, notification buttons, widget) and the iOS
native strings (permission prompts, alarm buttons) — with the English beside the
current translation, ordered so the screens that matter most come first.
Reviewers mark each line OK or Change and type a better wording; the hidden
"id" column maps every row back to its source for importing corrections.

Run from the repo root (needs openpyxl, e.g. in a throwaway venv):
    python3 -m venv /tmp/l10n-venv && /tmp/l10n-venv/bin/pip install openpyxl
    /tmp/l10n-venv/bin/python tool/l10n_review_sheets.py
Output: build/l10n_review/Rostrik translation review - <Language>.xlsx
"""

import html
import json
import re
from pathlib import Path

from openpyxl import Workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation

ROOT = Path(__file__).resolve().parent.parent
OUT = ROOT / 'build' / 'l10n_review'

LANGUAGES = {  # ARB code: (English name, native name, Android dir, iOS lproj)
    'es': ('Spanish', 'Español', 'values-es', 'es'),
    'pt': ('Portuguese (Brazil)', 'Português (Brasil)', 'values-pt', 'pt'),
    'de': ('German', 'Deutsch', 'values-de', 'de'),
    'fr': ('French', 'Français', 'values-fr', 'fr'),
    'it': ('Italian', 'Italiano', 'values-it', 'it'),
    'nl': ('Dutch', 'Nederlands', 'values-nl', 'nl'),
    'pl': ('Polish', 'Polski', 'values-pl', 'pl'),
    'tr': ('Turkish', 'Türkçe', 'values-tr', 'tr'),
    'id': ('Indonesian', 'Bahasa Indonesia', 'values-in', 'id'),
    'vi': ('Vietnamese', 'Tiếng Việt', 'values-vi', 'vi'),
    'ja': ('Japanese', '日本語', 'values-ja', 'ja'),
    'ko': ('Korean', '한국어', 'values-ko', 'ko'),
    'hi': ('Hindi', 'हिन्दी', 'values-hi', 'hi'),
    'ar': ('Arabic', 'العربية', 'values-ar', 'ar'),
}
RTL = {'ar'}

# Sections in review priority. Each claims strings by (source, key prefix);
# within a section, rows follow the order of these claims, then source order.
SECTIONS = [
    ('First launch',
     'The first screens a new user sees: the terms screen, the welcome screen and '
     'the permission requests. This is where people decide whether to trust the app.',
     [('arb', 'legal'), ('arb', 'welcome'), ('arb', 'perms'), ('arb', 'battery'), ('ios', 'NS')]),
    ('Free trial and purchase',
     'The free-trial notice, the screen asking people to buy once the trial ends, '
     'and the reminder before it ends.',
     [('arb', 'purchase'), ('arb', 'notifTrial')]),
    ('Alarm screen and notifications',
     'What the phone shows while an alarm is ringing, and the alarm notifications. '
     'Often read half-asleep: it must be instantly clear, and short enough to fit a button.',
     [('android', 'alarm_'), ('ios', 'Stop'), ('ios', 'Snooze'), ('arb', 'critical'),
      ('arb', 'notif'), ('android', 'channel_'), ('arb', 'sound'), ('arb', 'seed')]),
    ('Setting up a roster',
     'Choosing a shift pattern, building a custom roster, and importing a roster from a photo.',
     [('arb', 'roster'), ('arb', 'pattern'), ('arb', 'walkthrough'), ('arb', 'onb'),
      ('arb', 'builder'), ('arb', 'block'), ('arb', 'arm'), ('arb', 'draft'),
      ('arb', 'ai'), ('arb', 'ocr')]),
    ('Everyday screens',
     'The dashboard, timeline, alarms and sleep screens, and the roster tools.',
     [('arb', 'nav'), ('arb', 'dash'), ('arb', 'hero'), ('arb', 'alarms'), ('arb', 'create'),
      ('arb', 'timeline'), ('arb', 'filter'), ('arb', 'cal'), ('arb', 'day'), ('arb', 'shift'),
      ('arb', 'act'), ('arb', 'sleep'), ('android', 'sleep_'), ('arb', 'manage'),
      ('arb', 'mark'), ('arb', 'leave'), ('arb', 'work'), ('arb', 'tip')]),
    ('Settings',
     'The settings screen.',
     [('arb', 'settings'), ('android', 'ringtone_')]),
    ('Shared words and formats',
     'Buttons and words used on many screens, durations, weekday names, and the '
     'home-screen widget.',
     [('arb', 'common'), ('arb', 'duration'), ('arb', 'weekdays'), ('arb', 'app'),
      ('android', 'duration_'), ('android', 'widget_')]),
]

# Terms that must read the same everywhere; shown on the Start here sheet.
KEY_TERMS = [
    ('Shift (day / night)', 'arb', ['shiftTypeDay', 'shiftTypeNight']),
    ('Day off', 'arb', ['shiftTypeOff']),
    ('Alarms (tab)', 'arb', ['navAlarms']),
    ('Rotation', 'arb', ['dashRotation']),
    ('Critical shift', 'arb', ['createCriticalShift']),
    ('Snooze', 'android', ['alarm_action_snooze']),
    ('Dismiss', 'android', ['alarm_action_dismiss']),
    ('Annual leave', 'arb', ['leaveAnnual']),
    ('Sleep (tab)', 'arb', ['navSleep']),
]

SHORT_HINT_KEYS = re.compile(r'^(nav|filter|leave|sound|shiftType|common)|_action_|Short|Button|Tooltip')

INK = '161618'
ORANGE = 'FF8F00'
ORANGE_TINT = 'FFF1DC'
LINE = 'D9D6CF'
MUTED = '6B6A70'


# ---------------------------------------------------------------- sources

def read_arb(code):
    data = json.loads((ROOT / 'lib' / 'l10n' / f'app_{code}.arb').read_text(encoding='utf-8'))
    return {k: v for k, v in data.items() if not k.startswith('@')}


def read_android(dirname):
    text = (ROOT / 'android' / 'app' / 'src' / 'main' / 'res' / dirname / 'strings.xml').read_text(encoding='utf-8')
    out = {}
    for m in re.finditer(r'<string name="([a-z0-9_]+)"[^>]*>(.*?)</string>', text, re.S):
        v = html.unescape(m.group(2)).replace("\\'", "'").replace('\\"', '"').replace('\\n', '\n')
        out[m.group(1)] = v
    return out


def read_ios(lproj):
    out = {}
    for name in ('InfoPlist.strings', 'Localizable.strings'):
        path = ROOT / 'ios' / 'Runner' / f'{lproj}.lproj' / name
        if not path.exists():
            continue
        for m in re.finditer(r'"((?:[^"\\]|\\.)*)"\s*=\s*"((?:[^"\\]|\\.)*)"\s*;', path.read_text(encoding='utf-8')):
            out[m.group(1)] = m.group(2).replace('\\"', '"').replace('\\n', '\n')
    return out


# ---------------------------------------------------------------- display

def split_plural(message):
    """'{count, plural, one{A} other{B}}' -> [('one', 'A'), ('other', 'B')], else None."""
    m = re.match(r'^\{\s*(\w+)\s*,\s*plural\s*,(.*)\}$', message, re.S)
    if not m:
        return None
    body, cases, i = m.group(2), [], 0
    while i < len(body):
        cm = re.compile(r'\s*(=?\w+)\s*\{').match(body, i)
        if not cm:
            break
        depth, j = 1, cm.end()
        while j < len(body) and depth:
            depth += {'{': 1, '}': -1}.get(body[j], 0)
            j += 1
        cases.append((cm.group(1), body[cm.end():j - 1]))
        i = j
    return cases or None


def display(message):
    """Plural messages one form per line; line breaks shown as ↵."""
    cases = split_plural(message)
    if cases:
        return '\n'.join(f'[{case}]  {value}' for case, value in cases)
    return message.replace('\n', ' ↵\n')


def humanise(source, key):
    if source == 'ios':
        return {'Stop': 'iPhone alarm · Stop button', 'Snooze': 'iPhone alarm · Snooze button'}.get(
            key, 'iPhone permission prompt · ' + re.sub(r'^NS|UsageDescription$', '', key))
    if source == 'android':
        return 'Android · ' + key.replace('_', ' ')
    words = re.sub(r'([a-z0-9])([A-Z])', r'\1 \2', key).lower()
    return words


# ---------------------------------------------------------------- rows

def build_rows(en, lang):
    """[(section_index, [row dict])] in review order, every string exactly once."""
    claimed, sections = set(), []
    for si, (_, _, claims) in enumerate(SECTIONS):
        rows = []
        for source, prefix in claims:
            for key in en[source]:
                sid = (source, key)
                if sid in claimed or not key.startswith(prefix):
                    continue
                claimed.add(sid)
                rows.append({
                    'id': f'{source}:{key}',
                    'where': humanise(source, key),
                    'en': en[source][key],
                    'tr': lang[source].get(key, ''),
                    'short': bool(SHORT_HINT_KEYS.search(key)) and len(en[source][key]) <= 24,
                })
        sections.append((si, rows))
    missed = [f'{s}:{k}' for s in en for k in en[s] if (s, k) not in claimed]
    if missed:
        raise SystemExit(f'Strings not assigned to a section: {missed}')
    return sections


# ---------------------------------------------------------------- workbook

def thin(color=LINE):
    side = Side(style='thin', color=color)
    return Border(bottom=side)


def write_workbook(code, en):
    name, native, android_dir, lproj = LANGUAGES[code]
    lang = {'arb': read_arb(code), 'android': read_android(android_dir), 'ios': read_ios(lproj)}
    sections = build_rows(en, lang)
    total = sum(len(rows) for _, rows in sections)
    first_three = sum(len(rows) for si, rows in sections if si < 3)
    rtl = code in RTL

    wb = Workbook()
    start = wb.active
    start.title = 'Start here'
    review = wb.create_sheet('Review')

    # ---------------- Review sheet
    headers = ['Ref', 'Where it appears', 'English', f'{name} — current', 'Check',
               'Your suggested wording', 'Notes', 'id']
    widths = [7, 30, 50, 50, 11, 50, 30, 26]
    for col, (h, w) in enumerate(zip(headers, widths), start=1):
        cell = review.cell(row=1, column=col, value=h)
        cell.font = Font(bold=True, color='FFFFFF', size=11)
        cell.fill = PatternFill('solid', fgColor=INK)
        cell.alignment = Alignment(vertical='center', wrap_text=True)
        review.column_dimensions[cell.column_letter].width = w
    review.row_dimensions[1].height = 24
    review.column_dimensions['H'].hidden = True
    review.freeze_panes = 'C2'

    check = DataValidation(type='list', formula1='"OK,Change"', allow_blank=True,
                           showErrorMessage=True, errorTitle='Check',
                           error='Choose OK or Change.')
    review.add_data_validation(check)

    tr_align = Alignment(wrap_text=True, vertical='top',
                         horizontal='right' if rtl else 'left', readingOrder=2 if rtl else 1)
    top = Alignment(wrap_text=True, vertical='top')
    r = 2
    first_row = None
    for si, rows in sections:
        title, blurb, _ = SECTIONS[si]
        priority = '   ·   Review first' if si < 3 else ''
        review.merge_cells(start_row=r, start_column=1, end_row=r, end_column=7)
        head = review.cell(row=r, column=1, value=f'{si + 1}.  {title}{priority}\n{blurb}')
        head.font = Font(bold=True, size=12, color=INK)
        head.fill = PatternFill('solid', fgColor=ORANGE_TINT)
        head.alignment = Alignment(wrap_text=True, vertical='center')
        review.row_dimensions[r].height = 44
        r += 1
        for n, row in enumerate(rows, start=1):
            first_row = first_row or r
            review.cell(row=r, column=1, value=f'{si + 1}.{n:02d}').alignment = top
            where = review.cell(row=r, column=2, value=row['where'] + ('\n(short label — keep brief)' if row['short'] else ''))
            where.font = Font(color=MUTED, size=10)
            where.alignment = top
            review.cell(row=r, column=3, value=display(row['en'])).alignment = top
            review.cell(row=r, column=4, value=display(row['tr'])).alignment = tr_align
            chk = review.cell(row=r, column=5)
            chk.alignment = Alignment(vertical='top', horizontal='center')
            check.add(chk)
            review.cell(row=r, column=6).alignment = tr_align
            review.cell(row=r, column=7).alignment = top
            ident = review.cell(row=r, column=8, value=row['id'])
            ident.font = Font(color=MUTED, size=9)
            for col in range(1, 8):
                review.cell(row=r, column=col).border = thin()
            lines = max(display(row['en']).count('\n'), display(row['tr']).count('\n')) + 1
            longest = max(len(row['en']), len(row['tr']))
            review.row_dimensions[r].height = max(20, 15 * max(lines, -(-longest // 55)))
            r += 1
    last_row = r - 1

    rng = f'A{first_row}:G{last_row}'
    review.conditional_formatting.add(rng, FormulaRule(
        formula=[f'$E{first_row}="Change"'], fill=PatternFill('solid', fgColor='FDE4E1')))
    review.conditional_formatting.add(f'E{first_row}:E{last_row}', FormulaRule(
        formula=[f'$E{first_row}="OK"'], font=Font(color='1E7F4F', bold=True)))

    # ---------------- Start here sheet
    start.sheet_view.showGridLines = False
    start.column_dimensions['A'].width = 3
    start.column_dimensions['B'].width = 34
    start.column_dimensions['C'].width = 70

    def put(row, text, size=11, bold=False, color=INK, col=2, merge=True, height=None):
        if merge:
            start.merge_cells(start_row=row, start_column=col, end_row=row, end_column=3)
        c = start.cell(row=row, column=col, value=text)
        c.font = Font(size=size, bold=bold, color=color)
        c.alignment = Alignment(wrap_text=True, vertical='top')
        if height:
            start.row_dimensions[row].height = height
        return c

    put(2, 'Rostrik translation review', size=20, bold=True, height=30)
    put(3, f'{name}  ·  {native}', size=14, bold=True, color='B25E00', height=22)
    put(5, 'Rostrik is an alarm and roster app for shift workers. It turns a work roster into '
           'alarms that ring before each shift. Its screens were translated from English, and '
           'this file asks a native speaker to check they read naturally.', height=48)

    put(7, 'What to check', size=13, bold=True, height=20)
    checks = [
        ('Meaning', 'Does it say the same thing as the English?'),
        ('Natural wording', 'Would a native speaker say it this way? Word-for-word translation is not the goal.'),
        ('Tone', 'Friendly, plain and direct. People read it tired, sometimes at 3 am.'),
        ('Consistency', 'The same thing should have the same name everywhere — see Key terms below.'),
    ]
    row = 8
    for label, text in checks:
        put(row, label, bold=True, merge=False)
        put(row, text, col=3, merge=False, height=32)
        row += 1

    row += 1
    put(row, 'How to review', size=13, bold=True, height=20); row += 1
    steps = [
        'Open the Review sheet. The English is next to the current translation.',
        'In Check, choose OK if the line is fine, or Change if it needs fixing.',
        'For Change, type your wording in "Your suggested wording". Add a note if anything is unclear.',
        'Save the file and send it back. Partly finished is fine — anything you mark helps.',
    ]
    for i, text in enumerate(steps, start=1):
        put(row, f'{i}.  {text}', height=30); row += 1

    row += 1
    put(row, 'Short on time?', size=13, bold=True, height=20); row += 1
    put(row, f'Sections 1–3 matter most: first launch, the free trial and purchase screen, and the alarm '
             f'screen. That is {first_three} of the {total} lines.', height=32); row += 1

    row += 1
    put(row, 'Things to keep as they are', size=13, bold=True, height=20); row += 1
    rules = [
        ('{days}   {time}   %1$d', 'Filled in by the app with a number, time or name. Keep them unchanged, '
                                    'but move them wherever the sentence needs them.'),
        ('[one]   [other]   [few]', 'The same message for different numbers, e.g. "1 day" and "5 days". '
                                     'Check each line; your language may use more forms than English.'),
        ('↵', 'A line break on screen.'),
        ('(short label — keep brief)', 'A button, tab or label with little room. Short beats complete.'),
        ('Rostrik', 'The app\'s name. Never translated.'),
    ]
    for label, text in rules:
        put(row, label, bold=True, merge=False)
        put(row, text, col=3, merge=False, height=32)
        row += 1

    row += 1
    put(row, 'Key terms', size=13, bold=True, height=20); row += 1
    put(row, 'How these are translated now. If one should change, it probably changes everywhere — '
             'say so in a note.', height=32, color=MUTED); row += 1
    for label, source, keys in KEY_TERMS:
        values = ' / '.join(lang[source].get(k, '') for k in keys)
        put(row, label, bold=True, merge=False)
        c = put(row, values, col=3, merge=False, height=20)
        if rtl:
            c.alignment = Alignment(wrap_text=True, vertical='top', horizontal='right', readingOrder=2)
        row += 1

    row += 1
    put(row, 'Progress', size=13, bold=True, height=20); row += 1
    rng_e = f"Review!$E${first_row}:$E${last_row}"
    for label, formula in [
        ('Marked OK', f'=COUNTIF({rng_e},"OK")'),
        ('Marked Change', f'=COUNTIF({rng_e},"Change")'),
        ('Still to check', f'={total}-COUNTIF({rng_e},"OK")-COUNTIF({rng_e},"Change")'),
    ]:
        put(row, label, bold=True, merge=False)
        c = start.cell(row=row, column=3, value=formula)
        c.alignment = Alignment(horizontal='left')
        c.number_format = '0'
        row += 1

    bar = start.cell(row=1, column=2)
    start.merge_cells(start_row=1, start_column=2, end_row=1, end_column=3)
    bar.fill = PatternFill('solid', fgColor=ORANGE)
    start.row_dimensions[1].height = 6

    wb.active = 0
    OUT.mkdir(parents=True, exist_ok=True)
    path = OUT / f'Rostrik translation review - {name}.xlsx'
    wb.save(path)
    return path, total, first_three


def main():
    en = {'arb': read_arb('en'), 'android': read_android('values'), 'ios': read_ios('en')}
    for code in LANGUAGES:
        path, total, first = write_workbook(code, en)
        print(f'{path.relative_to(ROOT)}  ({total} lines, {first} in sections 1–3)')


if __name__ == '__main__':
    main()
